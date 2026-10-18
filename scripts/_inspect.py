import openpyxl, json
wb = openpyxl.load_workbook('Книги.xlsx', data_only=True)
ws = wb.active
print("sheet:", ws.title, "max_row:", ws.max_row, "max_col:", ws.max_column)
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    a, b, c = (list(r) + [None, None, None])[:3]
    if (a is None or str(a).strip() == '') and (b is None or str(b).strip() == ''):
        continue
    rows.append((a, b, c))
print("data rows:", len(rows))
for i, (a, b, c) in enumerate(rows, start=2):
    print(f"row{i} | {a!r} | {b!r} | {c!r}")
